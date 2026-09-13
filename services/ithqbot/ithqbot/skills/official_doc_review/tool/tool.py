"""Tool for reviewing OA/official documents from file_id or object storage."""

from __future__ import annotations

import asyncio
import json
import mimetypes
import os
import re
import tempfile
from collections import Counter
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import TYPE_CHECKING, Any

from docx import Document
from docx.oxml.ns import qn
from loguru import logger
from pydantic import BaseModel, Field

from ithqbot.agent.tools.base import Tool
from ithqbot.skills.official_doc_review.direct_handler import handle_direct_official_doc_review
from ithqbot.storage import build_storage_session, parse_storage_uri
from ithqbot.utils.helpers import extract_text_from_file_bytes

if TYPE_CHECKING:
    from ithqbot.agent.skills.base import SkillContext


_DEFAULT_DIMENSIONS = [
    "format",
    "typo",
    "political",
    "punctuation",
    "sequence",
    "redundancy",
    "terminology",
]
_CONTENT_DIMENSIONS = {"typo", "political", "punctuation", "redundancy", "terminology"}
_FORMAT_DIMENSIONS = {"format", "sequence"}
_CHINESE_NUMERALS = {
    "一": 1,
    "二": 2,
    "三": 3,
    "四": 4,
    "五": 5,
    "六": 6,
    "七": 7,
    "八": 8,
    "九": 9,
    "十": 10,
    "十一": 11,
    "十二": 12,
    "十三": 13,
    "十四": 14,
    "十五": 15,
    "十六": 16,
    "十七": 17,
    "十八": 18,
    "十九": 19,
    "二十": 20,
}
_CIRCLED_NUMERALS = {ch: idx for idx, ch in enumerate("①②③④⑤⑥⑦⑧⑨⑩", start=1)}
_TITLE_END_PUNCTUATION = "。；：，,.!?！？;:"
_BLACK_RGB_VALUES = {"000000", "00000000", "auto"}


@dataclass(slots=True)
class SourceDocument:
    source_ref: str
    filename: str
    mime: str
    content: str | bytes
    text: str


class ReviewProblemModel(BaseModel):
    dimension: str
    severity: str = "medium"
    location: str = ""
    description: str
    suggestion: str = ""


class ReviewLLMOutput(BaseModel):
    summary: str = ""
    problems: list[ReviewProblemModel] = Field(default_factory=list)


class OfficialDocReviewTool(Tool):
    _MAX_TEXT_CHARS = 12000

    def __init__(self, workspace: Path, config: Any, provider_factory: Any | None = None):
        super().__init__()
        self._workspace = workspace
        self._config = config
        self._provider_factory = provider_factory

    @property
    def name(self) -> str:
        return "official_doc_review"

    @property
    def timeout_s(self) -> int:
        return 240

    @property
    def description(self) -> str:
        return (
            "Review an uploaded OA/official document for format compliance, typos, sensitive wording, punctuation, "
            "numbering logic, redundancy, and terminology consistency."
        )

    @property
    def parameters(self) -> dict[str, Any]:
        return {
            "type": "object",
            "properties": {
                "file_id": {
                    "type": "string",
                    "description": "Platform file id returned by the file service.",
                },
                "storage_uri": {
                    "type": "string",
                    "description": "Object storage URI such as minio://bucket/path or s3://bucket/path.",
                },
                "rel_path": {
                    "type": "string",
                    "description": "Relative object path under the active default storage bucket.",
                },
                "file_name": {
                    "type": "string",
                    "description": "Optional display filename when storage_uri or rel_path does not contain a reliable suffix.",
                },
                "review_dimensions": {
                    "type": "array",
                    "items": {
                        "type": "string",
                        "enum": _DEFAULT_DIMENSIONS,
                    },
                    "description": "Optional subset of review dimensions.",
                },
                "generate_report": {
                    "type": "boolean",
                    "description": "Whether to save a markdown review report as a downloadable file.",
                },
            },
        }

    @property
    def direct_handler(self):
        return handle_direct_official_doc_review

    async def execute(
        self,
        file_id: str = "",
        storage_uri: str = "",
        rel_path: str = "",
        file_name: str = "",
        review_dimensions: list[str] | None = None,
        generate_report: bool = False,
        context: "SkillContext | None" = None,
        cancellation_token: Any | None = None,
        **kwargs: Any,
    ) -> str:
        self.throw_if_cancelled(cancellation_token)
        file_id = str(file_id or kwargs.get("file_id") or "").strip()
        storage_uri = str(storage_uri or kwargs.get("storage_uri") or "").strip()
        rel_path = str(rel_path or kwargs.get("rel_path") or "").strip()
        file_name = str(file_name or kwargs.get("file_name") or "").strip()
        generate_report = bool(generate_report or kwargs.get("generate_report", False))
        dimensions = self._normalize_dimensions(review_dimensions or kwargs.get("review_dimensions"))

        if not file_id and not storage_uri and not rel_path:
            return "错误：缺少必要参数，请至少提供 file_id、storage_uri 或 rel_path。"

        await self._emit_progress(context, 8, "reading", "正在读取待审查公文", step="load_source")

        try:
            source = await self._load_source(
                file_id=file_id,
                storage_uri=storage_uri,
                rel_path=rel_path,
                file_name=file_name,
                context=context,
            )
            self.throw_if_cancelled(cancellation_token)
        except Exception as exc:
            return f"错误：读取文档失败，{exc}"

        logger.info("[OfficialDocReview] Starting document review. Filename: {}, MIME: {}, Size: {} bytes", 
                    source.filename, source.mime, len(source.content))

        if not source.text.strip():
            return "错误：文档内容为空或暂不支持解析该格式。"

        await self._emit_progress(context, 30, "rules", "正在进行规则审查", step="rule_checks")

        problems = self._run_rule_checks(source, dimensions)

        await self._emit_progress(context, 58, "reasoning", "正在进行语义审查", step="llm_review")

        llm_output = await self._run_llm_review(source, dimensions, problems, context=context)
        self.throw_if_cancelled(cancellation_token)
        if llm_output is not None:
            problems.extend(self._normalize_problem(item.model_dump()) for item in llm_output.problems)

        problems = self._dedupe_problems(problems, dimensions)
        summary = self._build_summary(problems, llm_output.summary if llm_output else "")
        logger.info("[OfficialDocReview] Review completed. Total problems: {}", len(problems))
        message = f"审查完成，共发现 {len(problems)} 个问题。"

        payload: dict[str, Any] = {
            "status": "success",
            "message": message,
            "summary": summary,
            "source": {
                "file_id": file_id or "",
                "storage_uri": storage_uri or self._build_source_storage_uri(rel_path),
                "name": source.filename,
                "mime": source.mime,
            },
            "problems": problems,
            "review_dimensions": dimensions,
        }

        display_text, text_exceeded = self._build_display_text(
            source.filename,
            summary,
            problems,
            dimensions,
        )
        payload["llm_result"] = display_text

        should_generate_report = generate_report or len(problems) > 10 or text_exceeded
        if should_generate_report and context is not None:
            await self._emit_progress(context, 85, "report", "正在生成审查报告", step="save_report")
            report_text = self._build_report_markdown(source, summary, problems)
            report_filename = f"official_doc_review_{Path(source.filename).stem or 'report'}.md"
            logger.info("[OfficialDocReview] Stage: Report Generation. Filename: {}", report_filename)
            saved = await context.save_file(
                name=report_filename,
                content=report_text,
                mime="text/markdown",
            )
            file_meta = await context.get_file(saved["file_id"])
            download_url = f"/api/ext/icatmsg-client/files/download?path=/api/files/{saved['file_id']}/download"
            payload["files"] = [{
                "file_id": saved.get("file_id"),
                "name": file_meta.get("name"),
                "mime": file_meta.get("mime"),
                "size": file_meta.get("size"),
                "download_url": download_url,
                "storage": file_meta.get("storage", {}),
            }]
            if download_url:
                # We use the attachment list for the clickable link since relative markdown links are often blocked or fail to render
                payload["llm_result"] = display_text + f"\n\n📄 审查报告已生成，请在下方附件列表中点击下载。"
            else:
                payload["llm_result"] = display_text
            payload["report_file_id"] = saved.get("file_id")

        await self._emit_progress(context, 100, "done", "公文智审完成", step="finish")
        return json.dumps(payload, ensure_ascii=False)

    async def _emit_progress(
        self,
        context: "SkillContext | None",
        percent: int,
        stage: str,
        message: str,
        *,
        step: str,
    ) -> None:
        if context is None:
            return
        await context.emit_progress(
            percent,
            stage,
            message,
            progress_stage="skill_call",
            tool_name=self.name,
            call_type="skill",
            status_details={
                "execution": {
                    "step": step,
                    "stage": stage,
                }
            },
        )

    def _normalize_dimensions(self, value: Any) -> list[str]:
        if not isinstance(value, list) or not value:
            return list(_DEFAULT_DIMENSIONS)
        normalized: list[str] = []
        for item in value:
            text = str(item or "").strip().lower()
            if text in _DEFAULT_DIMENSIONS and text not in normalized:
                normalized.append(text)
        return normalized or list(_DEFAULT_DIMENSIONS)

    async def _load_source(
        self,
        *,
        file_id: str,
        storage_uri: str,
        rel_path: str,
        file_name: str,
        context: "SkillContext | None",
    ) -> SourceDocument:
        if file_id:
            if context is None:
                raise RuntimeError("通过 file_id 读取文件时缺少 skill context。")
            meta = await context.get_file(file_id)
            content = await context.read_file(file_id)
            filename = str(meta.get("name") or file_name or f"{file_id}.bin")
            mime = str(meta.get("mime") or self._guess_mime(filename))
            text = self._extract_text(content, filename)
            return SourceDocument(
                source_ref=file_id,
                filename=filename,
                mime=mime,
                content=content,
                text=text,
            )

        effective_storage_uri = storage_uri.strip() or self._build_source_storage_uri(rel_path)
        if not effective_storage_uri:
            raise RuntimeError("无法根据 rel_path 推断对象存储地址。")
        content, filename, mime = await self._read_from_storage(effective_storage_uri, file_name=file_name)
        return SourceDocument(
            source_ref=effective_storage_uri,
            filename=filename,
            mime=mime,
            content=content,
            text=self._extract_text(content, filename),
        )

    def _extract_text(self, content: str | bytes, filename: str) -> str:
        if isinstance(content, str):
            return content
        return extract_text_from_file_bytes(content, filename)

    async def _read_from_storage(self, storage_uri: str, *, file_name: str = "") -> tuple[bytes, str, str]:
        _, bucket, object_path = parse_storage_uri(storage_uri)
        if not bucket or not object_path:
            raise RuntimeError("storage_uri 不合法。")
        session = build_storage_session(self._config, bucket=bucket)
        fd, tmp_path = tempfile.mkstemp()
        os.close(fd)
        try:
            await asyncio.to_thread(session.client.fget_object, session.bucket, object_path, tmp_path)
            data = Path(tmp_path).read_bytes()
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)
        filename = file_name.strip() or Path(object_path).name or "document.bin"
        return data, filename, self._guess_mime(filename)

    def _guess_mime(self, filename: str) -> str:
        mime, _ = mimetypes.guess_type(filename)
        return mime or "application/octet-stream"

    def _build_source_storage_uri(self, rel_path: str) -> str:
        path_text = str(rel_path or "").strip().lstrip("/")
        if not path_text:
            return ""
        cfg = getattr(self._config, "get_active_storage_config", lambda: None)()
        bucket = str(getattr(cfg, "bucket", "") or "").strip()
        backend = str(getattr(cfg, "backend", "minio") or "minio").strip() or "minio"
        if not bucket:
            return ""
        return f"{backend}://{bucket}/{path_text}"

    def _run_rule_checks(self, source: SourceDocument, dimensions: list[str]) -> list[dict[str, Any]]:
        problems: list[dict[str, Any]] = []
        if "format" in dimensions and self._is_docx(source):
            problems.extend(self._inspect_docx_format(source.content, source.text))
        if "sequence" in dimensions:
            problems.extend(self._check_sequence_logic(source.text))
        if "redundancy" in dimensions:
            problems.extend(self._check_duplicate_paragraphs(source.text))
        if "format" in dimensions and len(source.text.strip()) > 3000:
            problems.append(
                self._problem(
                    "format",
                    "warning",
                    "全文",
                    f"正文长度约 {len(source.text.strip())} 字，超过“正文一般不超过3000字”的建议。",
                    "精简篇幅，保留主干信息，必要时拆分附件或附录。",
                )
            )
        return problems

    async def _run_llm_review(
        self,
        source: SourceDocument,
        dimensions: list[str],
        rule_problems: list[dict[str, Any]],
        *,
        context: "SkillContext | None",
    ) -> ReviewLLMOutput | None:
        if context is None:
            return None
        clipped_text = source.text.strip()[: self._MAX_TEXT_CHARS]
        rule_hints = self._summarize_rule_problems(rule_problems)
        dimension_desc = "、".join(dimensions)
        prompt = (
            "你是运营商文档资料审核员。请审查上传的 OA 文档，并只返回结构化 JSON。\n"
            "审查维度包括：\n"
            "1. format：是否符合 OA 文格式规范；\n"
            "2. typo：是否有错别字；\n"
            "3. political：是否有政治敏感词语或表述风险；\n"
            "4. punctuation：是否有标点符号错误；\n"
            "5. sequence：序号与层级是否存在逻辑错误；\n"
            "6. redundancy：是否存在语义重复、用词重复；\n"
            "7. terminology：专业术语前后是否不一致。\n\n"
            "OA 文格式规范要求：\n"
            "- 正文统一三号仿宋，一级标题三号黑体，二级标题三号楷体，三级标题三号仿宋；楷体和仿宋分别为“楷体_GB2312”“仿宋_GB2312”。\n"
            "- 正文表格字体为仿宋且建议小四，不要加粗加黑，不要出现彩色字或色块，正文字体颜色应统一为黑色。\n"
            "- 结构层次依次为“一、”“（一）”“1.”“（1）”“①”，原则上不超过五层；标题单独成行时不加标点；正文一般不超过3000字。\n"
            "- 段前段后为0行，单倍行距，每段前空2个字，不要使用首行缩进等自动套用格式。\n\n"
            f"本次重点审查维度：{dimension_desc}\n"
            f"文档文件名：{source.filename}\n"
            f"规则预审提示：{rule_hints or '无'}\n\n"
            "要求：\n"
            "- summary 用中文简洁概括总体结论；\n"
            "- problems 仅保留有依据的问题；\n"
            "- location 尽量写“第N段”“标题X”“表格第N行”等，不确定时可写“全文”；\n"
            "- suggestion 给出直接可执行的修改建议；\n"
            "- 不要重复规则预审中已经明确描述且无新增信息的问题。\n\n"
            f"待审文本如下：\n{clipped_text}"
        )
        logger.debug("[OfficialDocReview] Starting LLM review for {} dimensions", len(dimensions))
        try:
            route_profile = await context.resolve_llm_route_profile(
                task="extraction",
                prompt=prompt,
                max_tokens=1800,
            )
            logger.info("[OfficialDocReview] Stage: Semantic Analysis. Resolved Model: {} (Tier: {})", 
                        route_profile.active_model, route_profile.tier)
            result = await context.call_llm(
                task="extraction",
                prompt=prompt,
                output_model=ReviewLLMOutput,
                temperature=0.1,
                model=route_profile.active_model,
                max_tokens=route_profile.max_tokens,
                reasoning_effort=route_profile.reasoning_effort,
                retries=2,
            )
        except Exception as e:
            logger.warning("[OfficialDocReview] LLM review failed with exception: {}", e)
            return None
        
        if not isinstance(result, ReviewLLMOutput):
            logger.warning("[OfficialDocReview] LLM review returned unexpected type: {}", type(result))
            return None
            
        logger.debug("[OfficialDocReview] LLM review success, found {} problems", len(result.problems))
        return result

    def _inspect_docx_format(self, content: str | bytes, full_text: str) -> list[dict[str, Any]]:
        if isinstance(content, str):
            return []
        try:
            doc = Document(BytesIO(content))
        except Exception:
            return []
        problems: list[dict[str, Any]] = []
        headings = self._extract_heading_sequence(doc.paragraphs)
        for idx, paragraph in enumerate(doc.paragraphs, start=1):
            text = (paragraph.text or "").strip()
            if not text:
                continue
            heading_level = self._detect_heading_level(text)
            location = f"第{idx}段"
            if heading_level is not None and text[-1] in _TITLE_END_PUNCTUATION:
                problems.append(
                    self._problem(
                        "format",
                        "medium",
                        location,
                        f"标题“{text}”单独成行但末尾带有标点符号。",
                        "标题独立成行时去掉句号、分号、冒号等标点。",
                    )
                )
            problems.extend(self._check_paragraph_format(paragraph, text, heading_level, location))

        problems.extend(self._check_table_format(doc))
        if headings and max(item["level"] for item in headings) > 5:
            problems.append(
                self._problem(
                    "format",
                    "warning",
                    "全文",
                    "正文标题层级超过五层，不符合“原则上不要超过五层”的要求。",
                    "压缩层级，优先合并过深的小标题。",
                )
            )
        if len(full_text.strip()) > 3000:
            problems.append(
                self._problem(
                    "format",
                    "warning",
                    "全文",
                    f"正文长度约 {len(full_text.strip())} 字，超过公文建议字数。",
                    "适当精简背景铺陈、重复表述和非关键说明。",
                )
            )
        return problems

    def _check_paragraph_format(
        self,
        paragraph: Any,
        text: str,
        heading_level: int | None,
        location: str,
    ) -> list[dict[str, Any]]:
        problems: list[dict[str, Any]] = []
        expected_font = "仿宋_GB2312"
        if heading_level == 1:
            expected_font = "黑体"
        elif heading_level == 2:
            expected_font = "楷体_GB2312"
        elif heading_level == 3:
            expected_font = "仿宋_GB2312"

        observed_fonts = self._collect_paragraph_font_names(paragraph)
        for font_name in observed_fonts:
            if expected_font not in font_name:
                problems.append(
                    self._problem(
                        "format",
                        "medium",
                        location,
                        f"段落“{self._shorten(text)}”检测到字体“{font_name}”，与建议字体“{expected_font}”不一致。",
                        "按 OA 规范调整该段字体，正文用仿宋，一级标题用黑体，二级标题用楷体。",
                    )
                )
                break

        size_pt = self._resolve_paragraph_size_pt(paragraph)
        if size_pt is not None:
            if heading_level in {1, 2, 3, None} and abs(size_pt - 16) > 1.2 and heading_level is not None:
                problems.append(
                    self._problem(
                        "format",
                        "low",
                        location,
                        f"标题“{self._shorten(text)}”字号约为 {size_pt:.1f}pt，不符合三号字的常见设置。",
                        "将标题字号统一调整为三号。",
                    )
                )
            if heading_level is None and abs(size_pt - 16) > 1.2:
                problems.append(
                    self._problem(
                        "format",
                        "low",
                        location,
                        f"正文段落“{self._shorten(text)}”字号约为 {size_pt:.1f}pt，不符合三号仿宋的常见设置。",
                        "将正文统一调整为三号仿宋。",
                    )
                )

        para_format = paragraph.paragraph_format
        first_indent = getattr(para_format, "first_line_indent", None)
        if first_indent is not None:
            try:
                if abs(first_indent.pt) > 0.1:
                    problems.append(
                        self._problem(
                            "format",
                            "medium",
                            location,
                            f"段落“{self._shorten(text)}”使用了首行缩进等自动段落格式。",
                            "不要使用首行缩进；改为段首手动空 2 个字。",
                        )
                    )
            except Exception:
                first_indent = None

        for label, space in (("段前", para_format.space_before), ("段后", para_format.space_after)):
            if space is None:
                continue
            try:
                if abs(space.pt) > 0.1:
                    problems.append(
                        self._problem(
                            "format",
                            "low",
                            location,
                            f"段落“{self._shorten(text)}”{label}不为 0 行。",
                            "将正文段前、段后统一设置为 0 行。",
                        )
                    )
            except Exception:
                continue

        for run in paragraph.runs:
            run_text = (run.text or "").strip()
            if not run_text:
                continue
            if run.bold and heading_level is None:
                problems.append(
                    self._problem(
                        "format",
                        "medium",
                        location,
                        f"正文内容“{self._shorten(run_text)}”使用了加粗，不符合正文常规要求。",
                        "正文和正文表格内容尽量不用加粗加黑，保留统一字重。",
                    )
                )
                break
        color_issue = self._detect_color_issue(paragraph.runs)
        if color_issue:
            problems.append(
                self._problem(
                    "format",
                    "high",
                    location,
                    color_issue,
                    "将正文和标题颜色统一调整为一致的黑色，删除彩色字或色块。",
                )
            )
        return problems

    def _check_table_format(self, doc: Document) -> list[dict[str, Any]]:
        problems: list[dict[str, Any]] = []
        for table_idx, table in enumerate(doc.tables, start=1):
            for row_idx, row in enumerate(table.rows, start=1):
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        text = (paragraph.text or "").strip()
                        if not text:
                            continue
                        location = f"表格{table_idx}第{row_idx}行"
                        for font_name in self._collect_paragraph_font_names(paragraph):
                            if "仿宋" not in font_name:
                                problems.append(
                                    self._problem(
                                        "format",
                                        "medium",
                                        location,
                                        f"表格文字“{self._shorten(text)}”字体“{font_name}”不是仿宋字体。",
                                        "正文表格字体建议统一为仿宋_GB2312。",
                                    )
                                )
                                break
                        size_pt = self._resolve_paragraph_size_pt(paragraph)
                        if size_pt is not None and size_pt >= 15:
                            problems.append(
                                self._problem(
                                    "format",
                                    "low",
                                    location,
                                    f"表格文字“{self._shorten(text)}”字号约为 {size_pt:.1f}pt，偏大。",
                                    "正文表格字体建议小于三号，通常使用小四。",
                                )
                            )
                        if any(bool(run.bold) for run in paragraph.runs if (run.text or "").strip()):
                            problems.append(
                                self._problem(
                                    "format",
                                    "medium",
                                    location,
                                    f"表格文字“{self._shorten(text)}”存在加粗加黑。",
                                    "正文表格内容一般不要加粗加黑。",
                                )
                            )
                            break
                        color_issue = self._detect_color_issue(paragraph.runs)
                        if color_issue:
                            problems.append(
                                self._problem(
                                    "format",
                                    "high",
                                    location,
                                    color_issue,
                                    "表格中的文字颜色也应统一为黑色。",
                                )
                            )
        return problems

    def _collect_paragraph_font_names(self, paragraph: Any) -> list[str]:
        fonts: list[str] = []
        for run in paragraph.runs:
            fonts.extend(self._extract_run_font_names(run))
        style = getattr(paragraph, "style", None)
        if style is not None:
            style_font = getattr(style, "font", None)
            name = getattr(style_font, "name", None)
            if isinstance(name, str) and name.strip():
                fonts.append(name.strip())
        return self._unique_strings(fonts)

    def _extract_run_font_names(self, run: Any) -> list[str]:
        fonts: list[str] = []
        run_font = getattr(run, "font", None)
        font_name = getattr(run_font, "name", None)
        if isinstance(font_name, str) and font_name.strip():
            fonts.append(font_name.strip())
        rpr = getattr(getattr(run, "_element", None), "rPr", None)
        rfonts = getattr(rpr, "rFonts", None)
        if rfonts is not None:
            for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
                value = rfonts.get(qn(f"w:{attr}"))
                if isinstance(value, str) and value.strip():
                    fonts.append(value.strip())
        return self._unique_strings(fonts)

    def _resolve_paragraph_size_pt(self, paragraph: Any) -> float | None:
        for run in paragraph.runs:
            size = getattr(getattr(run, "font", None), "size", None)
            if size is not None:
                try:
                    return float(size.pt)
                except Exception:
                    continue
        style = getattr(paragraph, "style", None)
        size = getattr(getattr(style, "font", None), "size", None) if style is not None else None
        if size is None:
            return None
        try:
            return float(size.pt)
        except Exception:
            return None

    def _detect_color_issue(self, runs: list[Any]) -> str | None:
        colors: set[str] = set()
        for run in runs:
            rgb = getattr(getattr(getattr(run, "font", None), "color", None), "rgb", None)
            if rgb is None:
                continue
            value = str(rgb).strip().lower()
            if value:
                colors.add(value)
        if not colors:
            return None
        if any(color not in _BLACK_RGB_VALUES for color in colors):
            return "检测到正文或标题存在非黑色字体。"
        if len(colors) > 1:
            return "检测到文本颜色设置不一致。"
        return None

    def _check_sequence_logic(self, text: str) -> list[dict[str, Any]]:
        headings = self._extract_heading_sequence_from_text(text)
        if not headings:
            return []
        problems: list[dict[str, Any]] = []
        last_level = 0
        last_by_scope: dict[tuple[int, tuple[int, ...]], int] = {}
        parent_values: dict[int, int] = {}
        for item in headings:
            level = item["level"]
            value = item["value"]
            title = item["text"]
            line_no = item["line_no"]
            if last_level and level > last_level + 1:
                problems.append(
                    self._problem(
                        "sequence",
                        "medium",
                        f"第{line_no}行",
                        f"标题“{title}”的层级从第 {last_level} 层直接跳到了第 {level} 层。",
                        "检查标题层级，避免跨层跳号，按“一、/（一）/1./（1）/①”逐级展开。",
                    )
                )
            parent_key = tuple(parent_values.get(idx, 0) for idx in range(1, level))
            scope_key = (level, parent_key)
            expected = last_by_scope.get(scope_key, 0) + 1
            if value != expected:
                problems.append(
                    self._problem(
                        "sequence",
                        "high" if expected > 1 else "medium",
                        f"第{line_no}行",
                        f"标题“{title}”序号异常，当前为 {item['marker']}，按当前层级应为第 {expected} 项。",
                        "核对同层级标题是否漏号、跳号、重复编号，必要时重新顺排。",
                    )
                )
            last_by_scope[scope_key] = value
            parent_values[level] = value
            for deeper_level in range(level + 1, 10):
                parent_values.pop(deeper_level, None)
            last_level = level
        return problems

    def _extract_heading_sequence(self, paragraphs: list[Any]) -> list[dict[str, Any]]:
        items: list[dict[str, Any]] = []
        for idx, paragraph in enumerate(paragraphs, start=1):
            text = (paragraph.text or "").strip()
            marker = self._parse_heading_marker(text)
            if marker is None:
                continue
            items.append(
                {
                    "line_no": idx,
                    "text": text,
                    "level": marker["level"],
                    "value": marker["value"],
                    "marker": marker["marker"],
                }
            )
        return items

    def _extract_heading_sequence_from_text(self, text: str) -> list[dict[str, Any]]:
        items: list[dict[str, Any]] = []
        for idx, raw_line in enumerate(text.splitlines(), start=1):
            line = raw_line.strip()
            marker = self._parse_heading_marker(line)
            if marker is None:
                continue
            items.append(
                {
                    "line_no": idx,
                    "text": line,
                    "level": marker["level"],
                    "value": marker["value"],
                    "marker": marker["marker"],
                }
            )
        return items

    def _parse_heading_marker(self, text: str) -> dict[str, Any] | None:
        line = (text or "").strip()
        if not line:
            return None
        patterns = [
            (1, re.compile(r"^([一二三四五六七八九十]{1,3})、"), self._parse_cn_number),
            (2, re.compile(r"^（([一二三四五六七八九十]{1,3})）"), self._parse_cn_number),
            (3, re.compile(r"^(\d+)\."), lambda value: int(value)),
            (4, re.compile(r"^（(\d+)）"), lambda value: int(value)),
            (5, re.compile(r"^([①②③④⑤⑥⑦⑧⑨⑩])"), lambda value: _CIRCLED_NUMERALS.get(value, 0)),
        ]
        for level, pattern, parser in patterns:
            match = pattern.match(line)
            if not match:
                continue
            raw = match.group(1)
            value = parser(raw)
            if not value:
                return None
            return {
                "level": level,
                "value": value,
                "marker": match.group(0),
            }
        return None

    def _detect_heading_level(self, text: str) -> int | None:
        marker = self._parse_heading_marker(text)
        return None if marker is None else int(marker["level"])

    def _parse_cn_number(self, value: str) -> int:
        return _CHINESE_NUMERALS.get(str(value).strip(), 0)

    def _check_duplicate_paragraphs(self, text: str) -> list[dict[str, Any]]:
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        meaningful = [line for line in lines if len(line) >= 12]
        counter = Counter(meaningful)
        problems: list[dict[str, Any]] = []
        for line, count in counter.items():
            if count >= 2:
                problems.append(
                    self._problem(
                        "redundancy",
                        "medium",
                        "全文",
                        f"存在重复表述，共出现 {count} 次：“{self._shorten(line)}”。",
                        "合并重复句子，只保留一次准确表述。",
                    )
                )
        return problems

    def _summarize_rule_problems(self, problems: list[dict[str, Any]]) -> str:
        if not problems:
            return ""
        snippets: list[str] = []
        for item in problems[:8]:
            snippets.append(
                f"[{item.get('dimension')}/{item.get('location')}] {item.get('description')}"
            )
        return "；".join(snippets)

    def _build_summary(self, problems: list[dict[str, Any]], llm_summary: str) -> str:
        counts = Counter(str(item.get("dimension") or "other") for item in problems)
        ordered = [name for name in _DEFAULT_DIMENSIONS if counts.get(name)]
        parts = [f"{name} {counts[name]} 项" for name in ordered[:4]]
        suffix = f"重点问题涉及：{'、'.join(parts)}。" if parts else "未发现明确问题。"
        llm_text = llm_summary.strip()
        if llm_text:
            return f"{llm_text} {suffix}".strip()
        return suffix

    def _severity_badge(self, severity: str) -> str:
        badges = {
            "high": "🔴 高",
            "medium": "🟡 中",
            "warning": "⚠️ 提示",
            "low": "🔵 低",
        }
        return badges.get(severity, severity)

    def _is_content_issue(self, dimension: str) -> bool:
        return dimension in _CONTENT_DIMENSIONS

    def _build_display_text(
        self,
        source_name: str,
        summary: str,
        problems: list[dict[str, Any]],
        dimensions: list[str],
    ) -> tuple[str, bool]:
        severity_counts = Counter(str(item.get("severity") or "medium") for item in problems)
        dimension_counts = Counter(str(item.get("dimension") or "other") for item in problems)
        severity_parts = []
        for level, label in (("high", "高"), ("medium", "中"), ("warning", "提示"), ("low", "低")):
            if severity_counts.get(level):
                severity_parts.append(f"{label}{severity_counts[level]}项")
        dimension_parts = [
            f"{name} {dimension_counts[name]} 项"
            for name in dimensions
            if dimension_counts.get(name)
        ]

        lines = [
            "## 📋 公文智审报告",
            "",
            f"**审查文件：** {source_name}",
            f"**问题总数：** {len(problems)}",
            "",
            "### 📊 问题概览",
            "",
        ]
        if severity_parts:
            lines.append(f"严重程度：{' | '.join(severity_parts)}")
        if dimension_parts:
            lines.append(f"维度分布：{' | '.join(dimension_parts)}")

        lines.extend([
            "",
            "### 📝 总体结论",
            "",
            summary or "未发现明确问题。",
        ])

        if not problems:
            lines.extend(["", "## ✅ 问题明细", "", "未发现需要提示的问题。"])
            return "\n".join(lines), False

        content_issues = [p for p in problems if self._is_content_issue(p["dimension"])]
        format_issues = [p for p in problems if not self._is_content_issue(p["dimension"])]

        # 10-row logic: limit display to at most 10 problems in the message
        total_problems = len(problems)
        display_limit = 10
        shown_count = 0
        exceeded = total_problems > display_limit

        lines.append("")
        lines.append("## 🔍 问题明细")
        lines.append("")

        if content_issues:
            lines.append("### ⚠️ 内容问题")
            lines.append("")
            for idx, item in enumerate(content_issues, start=1):
                if shown_count >= display_limit:
                    break
                severity_marker = "**🔴 严重**" if item["severity"] == "high" else ""
                lines.append(f"**{idx}. {item['dimension']}** {self._severity_badge(item['severity'])} {severity_marker}")
                lines.append(f"- 📍 位置：{item['location'] or '全文'}")
                lines.append(f"- ❗ 问题：{item['description']}")
                lines.append(f"- 💡 建议：{item['suggestion'] or '请结合上下文人工复核并修订。'}")
                lines.append("")
                shown_count += 1

        if format_issues and shown_count < display_limit:
            lines.append("### 📐 格式问题")
            lines.append("")
            for idx, item in enumerate(format_issues, start=len(content_issues) + 1):
                if shown_count >= display_limit:
                    break
                severity_marker = "**🔴 严重**" if item["severity"] == "high" else ""
                lines.append(f"**{idx}. {item['dimension']}** {self._severity_badge(item['severity'])} {severity_marker}")
                lines.append(f"- 📍 位置：{item['location'] or '全文'}")
                lines.append(f"- ❗ 问题：{item['description']}")
                lines.append(f"- 💡 建议：{item['suggestion'] or '请结合上下文人工复核并修订。'}")
                lines.append("")
                shown_count += 1

        if exceeded:
            lines.append("---")
            lines.append("")
            lines.append(f"⚠️ 以上为前 {display_limit} 项问题预览，更多审查明细请下载完整报告查看。")

        return "\n".join(lines), exceeded

    def _build_report_markdown(
        self,
        source: SourceDocument,
        summary: str,
        problems: list[dict[str, Any]],
    ) -> str:
        severity_counts = Counter(str(item.get("severity") or "medium") for item in problems)
        severity_parts = []
        for level, label in (("high", "🔴 高"), ("medium", "🟡 中"), ("warning", "⚠️ 提示"), ("low", "🔵 低")):
            if severity_counts.get(level):
                severity_parts.append(f"{label}{severity_counts[level]}项")

        lines = [
            "# 📋 公文智审报告",
            "",
            f"**文档名称：** {source.filename}",
            f"**来源：** {source.source_ref}",
            f"**问题总数：** {len(problems)}",
            "",
            "## 📊 严重程度分布",
            "",
        ]
        if severity_parts:
            lines.append(" | ".join(severity_parts))
        else:
            lines.append("未发现问题")

        lines.extend([
            "",
            "## 📝 总体结论",
            "",
            summary or "未发现明确问题。",
        ])

        if not problems:
            lines.extend(["", "## ✅ 问题清单", "", "- 未发现需要提示的问题。"])
            return "\n".join(lines)

        content_issues = [p for p in problems if self._is_content_issue(p["dimension"])]
        format_issues = [p for p in problems if not self._is_content_issue(p["dimension"])]

        lines.append("")
        lines.append("## 🔍 问题清单")
        lines.append("")

        if content_issues:
            lines.append("### ⚠️ 内容问题")
            lines.append("")
            for idx, item in enumerate(content_issues, start=1):
                severity_marker = "**🔴 严重**" if item["severity"] == "high" else ""
                lines.append(f"**{idx}. {item['dimension']}** {self._severity_badge(item['severity'])} {severity_marker}")
                lines.append(f"- 📍 **位置：** {item['location'] or '全文'}")
                lines.append(f"- ❗ **问题：** {item['description']}")
                lines.append(f"- 💡 **建议：** {item['suggestion'] or '请结合上下文人工复核。'}")
                lines.append("")

        if format_issues:
            lines.append("### 📐 格式问题")
            lines.append("")
            for idx, item in enumerate(format_issues, start=len(content_issues) + 1):
                severity_marker = "**🔴 严重**" if item["severity"] == "high" else ""
                lines.append(f"**{idx}. {item['dimension']}** {self._severity_badge(item['severity'])} {severity_marker}")
                lines.append(f"- 📍 **位置：** {item['location'] or '全文'}")
                lines.append(f"- ❗ **问题：** {item['description']}")
                lines.append(f"- 💡 **建议：** {item['suggestion'] or '请结合上下文人工复核。'}")
                lines.append("")

        return "\n".join(lines).strip() + "\n"

    def _problem(
        self,
        dimension: str,
        severity: str,
        location: str,
        description: str,
        suggestion: str,
    ) -> dict[str, Any]:
        return {
            "dimension": dimension,
            "severity": severity,
            "location": location,
            "description": description,
            "suggestion": suggestion,
        }

    def _normalize_problem(self, raw: dict[str, Any]) -> dict[str, Any]:
        return {
            "dimension": str(raw.get("dimension") or "").strip().lower() or "format",
            "severity": str(raw.get("severity") or "medium").strip().lower() or "medium",
            "location": str(raw.get("location") or "").strip() or "全文",
            "description": str(raw.get("description") or "").strip(),
            "suggestion": str(raw.get("suggestion") or "").strip(),
        }

    def _dedupe_problems(self, problems: list[dict[str, Any]], dimensions: list[str]) -> list[dict[str, Any]]:
        normalized: list[dict[str, Any]] = []
        seen: set[tuple[str, str, str]] = set()
        for item in problems:
            normalized_item = self._normalize_problem(item)
            if not normalized_item["description"]:
                continue
            if normalized_item["dimension"] not in dimensions:
                continue
            key = (
                normalized_item["dimension"],
                normalized_item["location"],
                normalized_item["description"],
            )
            if key in seen:
                continue
            seen.add(key)
            normalized.append(normalized_item)
        severity_order = {"high": 0, "medium": 1, "warning": 2, "low": 3}
        normalized.sort(key=lambda item: (
            not self._is_content_issue(item["dimension"]),
            severity_order.get(item["severity"], 9),
            item["dimension"],
            item["location"],
        ))
        return normalized

    def _is_docx(self, source: SourceDocument) -> bool:
        return source.filename.lower().endswith(".docx") or "wordprocessingml.document" in source.mime.lower()

    def _shorten(self, text: str, limit: int = 28) -> str:
        compact = re.sub(r"\s+", " ", text or "").strip()
        if len(compact) <= limit:
            return compact
        return compact[: limit - 3] + "..."

    def _unique_strings(self, values: list[str]) -> list[str]:
        seen: set[str] = set()
        result: list[str] = []
        for item in values:
            text = str(item or "").strip()
            if not text or text in seen:
                continue
            seen.add(text)
            result.append(item)
        return result

    def _truncate_display_text(self, full_text: str, max_lines: int = 50) -> tuple[str, bool]:
        lines = full_text.splitlines()
        if len(lines) <= max_lines:
            return full_text, False
        truncated_lines = lines[:max_lines]
        truncated_lines.append("")
        truncated_lines.append("---")
        truncated_lines.append("")
        truncated_lines.append(f"⚠️ 以上为前 {max_lines} 行预览，更多审查明细请下载完整报告查看。")
        return "\n".join(truncated_lines), True
